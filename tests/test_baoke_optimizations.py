import asyncio
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
from sqlmodel import select

import app.db as db
import app.main as main
from app.engine.collection import KeywordCollector
from app.group_snapshots import ingest_group_snapshot
from app.models import (
    AccountActionTask,
    AccountOperationLog,
    DouyinAccount,
    GroupMember,
    GroupMemberEvent,
    KeywordCollectionJob,
    PublishTask,
    WebhookEvent,
    WebhookSubscription,
)
from app.reporting import build_keyword_collection_report


class BaokeOptimizationTests(unittest.TestCase):
    def setUp(self):
        self.previous_engine = db._engine
        self.tmp = tempfile.TemporaryDirectory()
        db.init_db(str(Path(self.tmp.name) / "optimizations.db"))
        with db.get_session() as session:
            account = DouyinAccount(
                platform="xhs", nickname="收藏账号", sec_uid="user-1",
                status="active", storage_state='{"cookies": []}')
            session.add(account); session.commit(); session.refresh(account)
            self.account_id = account.id

    def tearDown(self):
        if db._engine is not None:
            db._engine.dispose()
        db._engine = self.previous_engine
        self.tmp.cleanup()

    def test_collection_source_validation_preserves_keyword_compatibility(self):
        platform, keywords, _, _, options = main._validated_collection_input(
            main.KeywordCollectionIn(
                platform="douyin", account_id=self.account_id,
                keywords=["露营", "骑行"]))
        self.assertEqual((platform, keywords), ("douyin", ["露营", "骑行"]))
        self.assertEqual(options["source_type"], "keyword")

        _, keywords, _, _, options = main._validated_collection_input(
            main.KeywordCollectionIn(
                platform="xhs", account_id=self.account_id,
                source_type="favorite_board", source_ref="board-1",
                source_name="选题库"))
        self.assertEqual(keywords, [])
        self.assertEqual(options["source_ref"], "board-1")
        self.assertEqual(options["source_name"], "选题库")

        _, keywords, _, _, options = main._validated_collection_input(
            main.KeywordCollectionIn(
                platform="douyin", account_id=self.account_id,
                source_type="creator", source_ref="sec-user-1"))
        self.assertEqual(keywords, [])
        self.assertEqual(options["source_type"], "creator")
        self.assertEqual(options["source_ref"], "sec-user-1")

    def test_direct_favorite_source_paginates_and_preserves_token(self):
        cfg = SimpleNamespace(engine=SimpleNamespace(xhs_request_jitter=0))
        collector = KeywordCollector(cfg, SimpleNamespace(), SimpleNamespace())
        calls = []

        class Client:
            async def favorite_notes(self, **kwargs):
                calls.append(kwargs)
                if len(calls) == 1:
                    return {"notes": [{"note_id": "n1", "display_title": "收藏笔记",
                                        "xsec_token": "token-1"}],
                            "has_more": True, "cursor": "next"}
                return {"notes": [{"note_id": "n2"}], "has_more": False}

        job = KeywordCollectionJob(
            platform="xhs", account_id=self.account_id,
            source_type="favorite_board", source_ref="board-1",
            source_name="选题库", max_pages_per_keyword=3,
            max_contents_per_keyword=20)
        values, error = asyncio.run(
            collector._discover_xhs_favorites_api(Client(), job))
        self.assertEqual(error, "")
        self.assertEqual([value["note_id"] for value in values], ["n1", "n2"])
        self.assertEqual(values[0]["xsec_token"], "token-1")
        self.assertEqual(values[0]["xsec_source"], "pc_user")
        self.assertEqual(calls[0]["board_id"], "board-1")

    def test_group_snapshot_records_join_leave_and_profile_change(self):
        first = ingest_group_snapshot(
            platform="xhs", account_id=self.account_id, group_id="g1",
            group_name="测试群", members=[
                {"member_id": "u1", "nickname": "甲"},
                {"member_id": "u2", "nickname": "乙"},
            ])
        second = ingest_group_snapshot(
            platform="xhs", account_id=self.account_id, group_id="g1",
            group_name="测试群", members=[
                {"member_id": "u1", "nickname": "甲改"},
                {"member_id": "u3", "nickname": "丙"},
            ])
        self.assertEqual(first["joined"], ["u1", "u2"])
        self.assertEqual(second["joined"], ["u3"])
        self.assertEqual(second["left"], ["u2"])
        self.assertEqual(second["changed"], ["u1"])
        with db.get_session() as session:
            members = {row.member_id: row for row in session.exec(
                select(GroupMember).where(GroupMember.group_id == "g1")).all()}
            self.assertFalse(members["u2"].active)
            changes = session.exec(select(GroupMemberEvent).where(
                GroupMemberEvent.snapshot_id == second["snapshot_id"])).all()
            self.assertEqual({row.change_type for row in changes},
                             {"joined", "left", "changed"})
            self.assertGreaterEqual(len(session.exec(select(WebhookEvent)).all()), 2)

    def test_batch_stop_cancels_safe_tasks_and_skips_inflight_write(self):
        with db.get_session() as session:
            collection = KeywordCollectionJob(
                platform="xhs", account_id=self.account_id,
                source_type="favorite_all", source_name="全部收藏",
                status="running")
            publish = PublishTask(
                platform="xhs", account_id=self.account_id,
                status="publishing", title="正在发布")
            action = AccountActionTask(
                platform="xhs", account_id=self.account_id,
                status="pending", action="send_dm")
            session.add(collection); session.add(publish); session.add(action)
            session.commit(); session.refresh(collection); session.refresh(publish)
            session.refresh(action)
            ids = collection.id, publish.id, action.id
        result = asyncio.run(main.batch_stop_tasks(main.BatchStopIn(items=[
            main.BatchStopItem(queue_type="collections", id=ids[0]),
            main.BatchStopItem(queue_type="publishes", id=ids[1]),
            main.BatchStopItem(queue_type="actions", id=ids[2]),
        ])))
        self.assertEqual(result["stopped_count"], 2)
        self.assertEqual(result["skipped_count"], 1)
        with db.get_session() as session:
            collection = session.get(KeywordCollectionJob, ids[0])
            self.assertTrue(collection.cancel_requested)
            self.assertEqual(collection.status, "running")
            self.assertEqual(session.get(PublishTask, ids[1]).status, "publishing")
            self.assertEqual(session.get(AccountActionTask, ids[2]).status, "canceled")
            self.assertEqual(len(session.exec(select(AccountOperationLog)).all()), 2)
            self.assertGreaterEqual(len(session.exec(select(WebhookEvent)).all()), 2)

    def test_account_workbench_aggregates_tasks_and_never_returns_credentials(self):
        with db.get_session() as session:
            session.add(KeywordCollectionJob(
                platform="xhs", account_id=self.account_id,
                source_type="favorite_all", source_name="全部收藏",
                status="pending", blocked_reason="读取冷却"))
            session.add(AccountActionTask(
                platform="xhs", account_id=self.account_id,
                status="doing", action="send_dm"))
            session.add(PublishTask(
                platform="xhs", account_id=self.account_id,
                status="failed", title="失败样本"))
            session.commit()
        previous = main.open_browsers
        main.open_browsers = {}
        try:
            values = asyncio.run(main.accounts_workbench(platform="xhs"))
        finally:
            main.open_browsers = previous
        self.assertEqual(len(values), 1)
        self.assertEqual(values[0]["tasks"], {
            "total": 3, "pending": 0, "running": 1,
            "blocked": 1, "failed": 1,
        })
        self.assertNotIn("storage_state", values[0])
        self.assertNotIn("cookie", values[0])

    def test_webhook_read_contract_redacts_secret(self):
        with db.get_session() as session:
            row = WebhookSubscription(
                name="CRM", url="https://example.test/hook",
                secret="fixture-secret", event_types='["account.operation"]')
            session.add(row); session.commit(); session.refresh(row)
            payload = main._webhook_subscription_dict(row)
        self.assertTrue(payload["secret_configured"])
        self.assertNotIn("secret", payload)

    def test_collection_report_uses_non_keyword_source_metadata(self):
        job = SimpleNamespace(
            id=9, platform="xhs", source_type="favorite_board",
            source_ref="board-1", source_name="选题库", keywords="[]",
            status="completed")
        payload = build_keyword_collection_report(job, [], [])
        workbook = load_workbook(BytesIO(payload), read_only=True)
        summary, contents, comments, leads = workbook.worksheets
        self.assertEqual(workbook.properties.title, "CreatorHub 内容采集报告")
        summary_values = [cell.value for row in summary.iter_rows() for cell in row]
        self.assertIn("采集来源", summary_values)
        self.assertIn("选题库", summary_values)
        self.assertEqual(contents["B1"].value, "来源")
        self.assertEqual(comments["B1"].value, "来源")
        self.assertEqual(leads["A1"].value, "用户sec_uid")


if __name__ == "__main__":
    unittest.main()
